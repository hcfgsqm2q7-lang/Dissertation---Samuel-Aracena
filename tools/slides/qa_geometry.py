"""Geometric QA substitute: LibreOffice cannot render in this sandbox, so check
the defects a visual pass would catch (off-slide, margins, overlap, text overflow)
directly from the shape geometry."""
import math
import sys
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0
SW, SH = 13.333, 7.5
MARGIN = 0.5

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else "2026_08_10_slides.pptx")

def inches(v):
    return (v or 0) / EMU

def est_lines(text, width_in, pt):
    """Rough wrap estimate: avg glyph advance ~0.5em for Calibri/Cambria mixed case."""
    if not text.strip():
        return 0
    cpl = max(1, int((width_in * 72.0) / (pt * 0.50)))
    n = 0
    for para in text.split("\n"):
        n += max(1, math.ceil(len(para) / cpl))
    return n

def max_pt(shape):
    pts = []
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                pts.append(r.font.size.pt)
    return max(pts) if pts else 14.0

issues = []
for idx, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        x, y = inches(sh.left), inches(sh.top)
        w, h = inches(sh.width), inches(sh.height)
        name = sh.shape_type
        txt = sh.text_frame.text if sh.has_text_frame else ""

        # 1. off-slide / margin. Decorative ellipses are intentionally bled off-edge.
        decorative = not txt.strip()   # all deliberate bleed shapes carry no text
        if not decorative:
            if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
                issues.append(f"S{idx} OFF-SLIDE  {str(name)[:18]:18s} x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f} :: {txt[:40]!r}")
            elif txt.strip() and (x < MARGIN - 0.01 or y < MARGIN - 0.01 or x + w > SW - MARGIN + 0.01 or y + h > SH - MARGIN + 0.01):
                issues.append(f"S{idx} MARGIN<0.5 x={x:.2f} y={y:.2f} r={x+w:.2f} b={y+h:.2f} :: {txt[:40]!r}")

        # 2. text overflow
        if txt.strip():
            pt = max_pt(sh)
            lines = est_lines(txt, w, pt)
            need = lines * pt * 1.22 / 72.0
            if need > h + 0.06:
                issues.append(f"S{idx} OVERFLOW?  need={need:.2f}in have={h:.2f}in pt={pt:.0f} lines={lines} :: {txt[:45]!r}")
            boxes.append((x, y, w, h, txt))

    # 3. overlap between text-bearing boxes
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[i]
            bx, by, bw, bh, bt = boxes[j]
            ox = min(ax + aw, bx + bw) - max(ax, bx)
            oy = min(ay + ah, by + bh) - max(ay, by)
            if ox > 0.12 and oy > 0.12:
                issues.append(f"S{idx} OVERLAP    {ox:.2f}x{oy:.2f}in :: {at[:26]!r} vs {bt[:26]!r}")

print(f"slides: {len(prs.slides)}   issues: {len(issues)}")
for i in issues:
    print(" ", i)
